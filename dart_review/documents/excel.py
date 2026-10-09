"""The workbook: the statements as tables, where every value came from, and what was checked.

Sheets:
- 안내: how to read the workbook (all amounts in 원; yellow cells need checking against the source).
- 재무제표: each statement as on the page, values in 원, cells marked "확인 필요" filled yellow with a comment.
- 셀 출처: one line per value: page, position, printed text, unit, method (text layer / OCR), confidence, status.
- 검증: the checks that failed and the problems of each statement.
"""

from __future__ import annotations

from pathlib import Path

from .statement import Statement
from .verify import REVIEW, CellCheck

KINDS = {"BS": "재무상태표", "IS": "손익계산서", "CIS": "포괄손익계산서", "CF": "현금흐름표", "SCE": "자본변동표"}


METHODS = {"text": "PDF 글자", "ocr": "OCR", "ai": "AI"}


def write_workbook(path: Path, statements: list[Statement], checks: list[CellCheck], source_name: str,
                   problems: list[str] | None = None, reader: str = "rules") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed: pip install -e '.[docs]'") from exc
    flagged = {(c.statement, c.row, c.column): c for c in checks}
    yellow = PatternFill("solid", fgColor="FFF2CC")
    bold = Font(bold=True)
    book = Workbook()

    guide = book.active
    guide.title = "안내"
    review = sum(c.status == REVIEW for c in checks)
    for line in (f"원본: {source_name}", "모든 금액은 원 단위로 바꿔 적었다 (원본 단위는 '셀 출처' 시트).",
                 f"검사한 칸 {len(checks)}개 가운데 '확인 필요' {review}개: 노란 칸은 원본과 대조해 확인한다.",
                 "숫자는 원본에서만 가져왔고, 의심되는 칸도 고치지 않았다." if reader != "ai" else
                 "AI(로컬 비전 모델)가 표를 읽었고, 각 숫자를 원본(PDF 글자 또는 OCR)의 같은 줄과 대조했다. "
                 "원본에서 같은 숫자를 찾지 못한 칸은 '확인 필요'다. AI가 읽은 숫자를 원본으로 고치지는 않았다."):
        guide.append([line])

    sheet = book.create_sheet("재무제표")
    for s_index, statement in enumerate(statements):
        sheet.append([f"{statement.title} (p.{statement.page}, 단위 원본: {statement.unit_text or '알 수 없음'})"])
        sheet.cell(sheet.max_row, 1).font = bold
        sheet.append(["계정", "표준 계정"] + [_header(column) for column in statement.columns])
        for cell in sheet[sheet.max_row]:
            cell.font = bold
        for r_index, row in enumerate(statement.rows):
            values = [row.amounts[c].value if c in row.amounts else None for c in range(len(statement.columns))]
            sheet.append([row.label, row.account or ""] + values)
            for c_index in row.amounts:
                target = sheet.cell(sheet.max_row, 3 + c_index)
                target.number_format = "#,##0"
                check = flagged.get((s_index, r_index, c_index))
                if check and check.status == REVIEW:
                    target.fill = yellow
                    target.comment = Comment("확인 필요: " + "; ".join(check.reasons), "dart_review")
        sheet.append([])

    sources = book.create_sheet("셀 출처")
    sources.append(["재무제표", "계정", "기간 열", "금액(원)", "원문 표기", "원본 단위", "페이지", "x0", "top", "x1", "bottom",
                    "읽은 방법", "증거", "OCR 확신도", "상태", "이유"])
    for s_index, statement in enumerate(statements):
        for r_index, row in enumerate(statement.rows):
            for c_index, amount in row.amounts.items():
                check = flagged.get((s_index, r_index, c_index))
                cell = amount.cell
                sources.append([statement.title, row.label, _header(statement.columns[c_index]), amount.value, amount.printed,
                                statement.unit_text, cell.page, round(cell.x0, 1), round(cell.top, 1), round(cell.x1, 1),
                                round(cell.bottom, 1), _method(cell, check, reader), _evidence(check, reader),
                                cell.confidence, check.status if check else "", "; ".join(check.reasons) if check else ""])
                sources.cell(sources.max_row, 4).number_format = "#,##0"

    sheet_problems = book.create_sheet("검증")
    sheet_problems.append(["재무제표", "계정", "기간 열", "원문 표기", "이유"])
    for check in checks:
        if check.status != REVIEW:
            continue
        statement = statements[check.statement]
        row = statement.rows[check.row]
        sheet_problems.append([statement.title, row.label, _header(statement.columns[check.column]),
                               row.amounts[check.column].printed, "; ".join(check.reasons)])
    for statement in statements:
        for problem in statement.problems:
            sheet_problems.append([statement.title, "", "", "", problem])
    for problem in problems or []:
        sheet_problems.append(["(문서 전체)", "", "", "", problem])
    book.save(path)


def _method(cell, check, reader: str) -> str:
    """How the value was read: by the AI (its evidence is in the next column) or straight from the page."""
    return "AI" if reader == "ai" else METHODS.get(cell.method, cell.method)


def _evidence(check, reader: str) -> str:
    if reader != "ai":
        return ""
    return check.evidence if check and check.evidence else "없음"


def _header(column) -> str:
    parts = [column.header or "(머리글 없음)"]
    if column.end:
        parts.append(f"기말 {column.end}")
    elif column.year:
        parts.append(f"{column.year}년")
    return " / ".join(parts)
