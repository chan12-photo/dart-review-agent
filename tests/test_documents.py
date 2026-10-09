"""Documents to verified spreadsheets (dart_review.documents).

The table, statement, and check logic run on synthetic words without any optional package. The PDF, OCR,
and Excel paths run on sample PDFs whose every value is known (samples.py) and are skipped when their
optional dependency is missing (pip install -e '.[docs]', brew install tesseract tesseract-lang).
"""

from importlib.util import find_spec
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from dart_review.documents.samples import SampleStatement, korean_font, scan_pdf, text_pdf  # noqa: E402
from dart_review.documents.statement import account_of, clean_label, find_statements, parse_amount  # noqa: E402
from dart_review.documents.table import group_lines  # noqa: E402
from dart_review.documents.verify import REVIEW, check, separator_problem, uses_separators  # noqa: E402
from dart_review.documents.words import Word  # noqa: E402

HAS_PDF = find_spec("pdfplumber") is not None and find_spec("pypdfium2") is not None
HAS_EXCEL = find_spec("openpyxl") is not None
HAS_OCR = HAS_PDF and shutil.which("tesseract") is not None and korean_font() is not None

BALANCE = SampleStatement(
    "재무상태표", "주식회사 예시", ["제 56 기 2025년 12월 31일 현재", "제 55 기 2024년 12월 31일 현재"], "백만원",
    ["제 56 (당) 기", "제 55 (전) 기"],
    [("자산", ["", ""]), ("Ⅰ. 유동자산", ["227,062,266", "202,945,624"]), ("Ⅱ. 비유동자산", ["296,597,320", "288,361,693"]),
     ("자산총계", ["523,659,586", "491,307,317"]), ("부채총계", ["112,339,878", "102,231,416"]),
     ("자본총계", ["411,319,708", "389,075,901"])])
INCOME = SampleStatement(
    "손익계산서", "주식회사 예시", ["제 56 기 2025년 1월 1일부터 2025년 12월 31일까지", "제 55 기 2024년 1월 1일부터 2024년 12월 31일까지"],
    "천원", ["제 56 (당) 기", "제 55 (전) 기"],
    [("매출액", ["333,605,938", "300,870,903"]), ("영업이익", ["43,601,051", "-"]), ("당기순이익(손실)", ["(1,234,567)", "34,451,351"])])


def expected(statement: SampleStatement, unit: int) -> dict[str, list[int | None]]:
    return {clean_label(label): [parse_amount(text, unit) for text in amounts]
            for label, amounts in statement.rows if any(amounts)}


def words_for(rows: list[tuple[float, list[tuple[float, float, str]]]], page: int = 1) -> list[Word]:
    """Synthetic words: rows of (top, [(x0, x1, text)])."""
    return [Word(text, page, x0, x1, top, top + 9) for top, items in rows for x0, x1, text in items]


class AmountTests(unittest.TestCase):
    def test_amounts_as_printed(self):
        self.assertEqual(parse_amount("1,234", 1_000), 1_234_000)
        for printed in ("(1,234)", "△1,234", "-1,234", "（1,234）"):
            self.assertEqual(parse_amount(printed, 1), -1234, printed)
        self.assertIsNone(parse_amount("-", 1))
        with self.assertRaises(ValueError):
            parse_amount("주석", 1)

    def test_separators(self):
        self.assertIsNone(separator_problem("300,870,903"))
        self.assertIsNone(separator_problem("(1,234,567)"))
        self.assertIn("마침표", separator_problem("300,870.903"))  # the misread seen in the first OCR test
        self.assertIn("세 자리", separator_problem("3,00,870"))

    def test_labels(self):
        self.assertEqual(clean_label("Ⅰ. 유동자산"), "유동자산")
        self.assertEqual(clean_label("| . 유 동 자산", ocr=True), "유동자산")
        self.assertEqual(clean_label("현금및현금성자산 (주4,5)"), "현금및현금성자산")
        self.assertEqual(clean_label("© 비 유 동 자산", ocr=True), "비유동자산")
        self.assertEqual(clean_label("i. 비 유 동 자산", ocr=True), "비유동자산")  # OCR read "Ⅱ." as "i."
        self.assertEqual(account_of("당기순이익(손실)"), "net_income")
        self.assertEqual(account_of("자 산 총 계"), "total_assets")
        self.assertIsNone(account_of("자본총계"))


class TableTests(unittest.TestCase):
    ROWS = [(40, [(250, 300, "재무상태표")]),
            (60, [(250, 330, "제 56 기 2025년 12월 31일 현재")]),
            (80, [(470, 540, "(단위: 원)")]),
            (100, [(60, 90, "과 목"), (330, 380, "제 56 기"), (470, 520, "제 55 기")]),
            (120, [(60, 100, "자산총계"), (320, 380, "1,000"), (480, 520, "900")]),
            (140, [(60, 100, "부채총계"), (330, 380, "400"), (480, 520, "300")]),
            (160, [(60, 100, "자본총계"), (330, 380, "600"), (480, 520, "650")])]

    def test_a_statement_from_positions(self):
        [statement] = find_statements(group_lines(words_for(self.ROWS)))
        self.assertEqual((statement.kind, statement.unit), ("BS", 1))
        self.assertEqual([column.header for column in statement.columns], ["제 56 기", "제 55 기"])
        self.assertEqual(statement.columns[0].end, "2025-12-31")
        self.assertEqual([row.label for row in statement.rows], ["자산총계", "부채총계", "자본총계"])
        self.assertEqual(statement.rows[0].amounts[0].value, 1000)

    def test_the_balance_identity_marks_all_three_cells(self):
        [statement] = find_statements(group_lines(words_for(self.ROWS)))
        flagged = {(item.row, item.column) for item in check([statement]) if item.status == REVIEW}
        # 2025: 1,000 = 400 + 600 holds; 2024: 900 != 300 + 650, so its three cells need review
        self.assertEqual(flagged, {(0, 1), (1, 1), (2, 1)})

    def test_a_missing_unit_needs_review(self):
        rows = [row for row in self.ROWS if "단위" not in row[1][0][2]]
        [statement] = find_statements(group_lines(words_for(rows)))
        self.assertIsNone(statement.unit)
        self.assertTrue(all(item.status == REVIEW for item in check([statement])))

    def test_a_missing_separator_needs_review_where_the_document_uses_them(self):
        rows = [(40, [(250, 300, "손익계산서")]), (80, [(470, 540, "(단위: 원)")]),
                (100, [(60, 90, "과 목"), (330, 380, "당기"), (470, 520, "전기")]),
                (120, [(60, 100, "매출액"), (320, 380, "333,605,938"), (460, 520, "300870903")]),
                (140, [(60, 100, "영업이익"), (320, 380, "43,601,051"), (460, 520, "32,725,961")]),
                (160, [(60, 100, "법인세비용"), (330, 380, "8,422,019"), (470, 520, "7,140,108")]),
                (180, [(60, 100, "당기순이익"), (320, 380, "34,451,351"), (460, 520, "23,567,302")])]
        [statement] = find_statements(group_lines(words_for(rows)))
        self.assertTrue(uses_separators([amount for row in statement.rows for amount in row.amounts.values()]))
        flagged = [(item.row, item.column, item.reasons) for item in check([statement]) if item.status == REVIEW]
        self.assertEqual([(row, column) for row, column, _ in flagged], [(0, 1)])


class OcrDamageTests(unittest.TestCase):
    """What a degraded scan did in the stress test, on synthetic OCR words."""

    def ocr_words(self, rows, confidence=95.0):
        return [Word(text, 1, x0, x1, top, top + 9, "ocr", confidence, (1, 1, number))
                for number, (top, items) in enumerate(rows) for x0, x1, text in items]

    def test_a_number_split_after_a_separator_is_joined(self):
        rows = [(40, [(250, 300, "재무상태표")]), (80, [(470, 540, "(단위: 원)")]),
                (100, [(60, 90, "과 목"), (330, 380, "당기"), (470, 520, "전기")]),
                (120, [(60, 100, "비유동자산"), (320, 380, "296,597,320"), (455, 470, "288,"), (472, 520, "361,693")]),
                (140, [(60, 100, "자산총계"), (320, 380, "523,659,586"), (460, 520, "491,307,317")])]
        [statement] = find_statements(group_lines(self.ocr_words(rows)))
        self.assertEqual(statement.rows[0].amounts[1].printed, "288,361,693")
        self.assertEqual(statement.rows[0].amounts[1].value, 288_361_693)

    def test_an_unplaced_number_marks_its_row(self):
        rows = [(40, [(250, 300, "손익계산서")]), (80, [(470, 540, "(단위: 원)")]),
                (100, [(60, 90, "과 목"), (330, 380, "당기"), (470, 520, "전기")]),
                (120, [(60, 100, "매출액"), (320, 380, "333,605,938"), (400, 430, "777,000"), (460, 520, "300,870,903")]),
                (140, [(60, 100, "영업이익"), (320, 380, "43,601,051"), (460, 520, "32,725,961")]),
                (160, [(60, 100, "당기순이익"), (320, 380, "34,451,351"), (460, 520, "23,567,302")])]
        [statement] = find_statements(group_lines(self.ocr_words(rows)))
        self.assertEqual(statement.rows[0].unplaced, 1)
        flagged = {(item.row, item.column) for item in check([statement]) if item.status == REVIEW}
        self.assertEqual(flagged, {(0, 0), (0, 1)})
        self.assertTrue(any("표에 넣지 못한 금액 1개" in problem for problem in statement.problems))

    def test_a_balance_sheet_without_its_totals_says_it_was_not_checked(self):
        rows = [(40, [(250, 300, "재무상태표")]), (80, [(470, 540, "(단위: 원)")]),
                (100, [(60, 90, "과 목"), (330, 380, "당기"), (470, 520, "전기")]),
                (120, [(60, 100, "자산총계"), (330, 380, "1,000"), (480, 520, "900")]),
                (140, [(60, 100, "부채출계"), (330, 380, "400"), (480, 520, "300")])]  # "총" misread
        [statement] = find_statements(group_lines(self.ocr_words(rows)))
        self.assertIn("자산총계·부채총계·자본총계 줄을 모두 찾지 못해 합계 검산을 하지 못했다", statement.problems)

    def test_an_empty_cell_is_reported(self):
        rows = [(40, [(250, 300, "손익계산서")]), (80, [(470, 540, "(단위: 원)")]),
                (100, [(60, 90, "과 목"), (330, 380, "당기"), (470, 520, "전기")]),
                (120, [(60, 100, "매출액"), (320, 380, "333,605,938"), (460, 520, "300,870,903")]),
                (140, [(60, 100, "영업이익"), (320, 380, "43,601,051")]),
                (160, [(60, 100, "당기순이익"), (320, 380, "34,451,351"), (460, 520, "23,567,302")])]
        [statement] = find_statements(group_lines(self.ocr_words(rows)))
        self.assertTrue(any(problem.startswith("빈 칸 1개 (영업이익)") for problem in statement.problems), statement.problems)


@unittest.skipUnless(HAS_PDF, "pip install -e '.[docs]'")
class TextPdfTests(unittest.TestCase):
    def test_every_value_of_a_text_pdf(self):
        from dart_review.documents import extract
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.pdf"
            text_pdf(path, [BALANCE, INCOME])
            result = extract(path, ocr="never")
        self.assertEqual(result.pages_by_method, {"text": [1, 2]})
        balance, income = result.statements
        self.assertEqual({row.label: [row.amounts[c].value for c in sorted(row.amounts)] for row in balance.rows},
                         expected(BALANCE, 1_000_000))
        self.assertEqual({row.label: [row.amounts[c].value for c in sorted(row.amounts)] for row in income.rows},
                         expected(INCOME, 1_000))
        self.assertEqual([(c.year, c.end) for c in balance.columns], [(2025, "2025-12-31"), (2024, "2024-12-31")])
        self.assertEqual([c.header for c in balance.columns], ["제 56 (당) 기", "제 55 (전) 기"])
        self.assertEqual(result.needs_review, 0)


@unittest.skipUnless(HAS_OCR, "needs pdfplumber, pypdfium2, tesseract with kor, and a Korean font")
class ScanTests(unittest.TestCase):
    def test_a_scan_is_read_by_ocr_and_checked(self):
        from dart_review.documents import extract
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "scan.pdf"
            scan_pdf(path, [BALANCE, INCOME])
            result = extract(path, ocr="auto")
        self.assertEqual(result.pages_by_method, {"ocr": [1, 2]})
        balance, income = result.statements
        wrong, unread = [], []
        for statement, sample, unit in ((balance, BALANCE, 1_000_000), (income, INCOME, 1_000)):
            rows = {row.label: row for row in statement.rows}
            for label, values in expected(sample, unit).items():
                for column, value in enumerate(values):
                    amount = rows[label].amounts.get(column) if label in rows else None
                    if amount is None:
                        unread.append((statement, label))
                    elif amount.value != value:
                        wrong.append((label, amount.printed, value))
        # no wrong value on a clean scan; a cell OCR could not read (a lone "-" may vanish) must be reported
        self.assertEqual(wrong, [])
        for statement, label in unread:
            self.assertTrue(any(label in problem for problem in statement.problems), (label, statement.problems))

    def test_a_scan_too_poor_to_read_says_so(self):
        from dart_review.documents import extract
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "poor.pdf"
            scan_pdf(path, [BALANCE], dpi=90, rotate=2.5, blur=1.6)
            result = extract(path, ocr="auto")
        self.assertTrue(any("스캔 품질이 낮다" in problem for problem in result.problems), result.problems)


@unittest.skipUnless(HAS_PDF and HAS_EXCEL, "pip install -e '.[docs]'")
class WorkbookTests(unittest.TestCase):
    def test_sheets_values_and_marked_cells(self):
        from openpyxl import load_workbook
        from dart_review.documents import extract
        rows = list(BALANCE.rows)
        rows[-1] = ("자본총계", ["411,319,709", "389,075,901"])  # breaks the 2025 balance identity on purpose
        broken = SampleStatement(BALANCE.title, BALANCE.company, BALANCE.period_lines, BALANCE.unit, BALANCE.headers, rows)
        with tempfile.TemporaryDirectory() as folder:
            path, out = Path(folder) / "sample.pdf", Path(folder) / "sample.xlsx"
            text_pdf(path, [broken])
            result = extract(path, out, ocr="never")
            book = load_workbook(out)
        self.assertEqual(book.sheetnames, ["안내", "재무제표", "셀 출처", "검증"])
        self.assertEqual(result.needs_review, 3)
        sheet = book["재무제표"]
        marked = [cell for row in sheet.iter_rows() for cell in row if cell.comment is not None]
        self.assertEqual(len(marked), 3)
        self.assertTrue(all("자산총계 = 부채총계 + 자본총계" in cell.comment.text for cell in marked))
        values = {row[0].value: row[2].value for row in sheet.iter_rows(min_row=3) if row[0].value}
        self.assertEqual(values["자산총계"], 523_659_586_000_000)
        self.assertEqual(book["검증"].max_row, 4)  # header + three cells


if __name__ == "__main__":
    unittest.main()
